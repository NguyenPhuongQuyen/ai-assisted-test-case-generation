import os, fitz, csv, openpyxl

from werkzeug.utils import secure_filename
from config import Config
from docx import Document

def allowed_file(filename):
    return '.' in filename and \
        filename.rsplit('.', 1)[1].lower() in Config.ALLOWED_EXTENSIONS

def save_uploaded_file(file):
    filename = secure_filename(file.filename)
    os.makedirs(Config.UPLOAD_FOLDER, exist_ok=True)
    file_path = os.path.join(Config.UPLOAD_FOLDER, filename)
    file.save(file_path)
    return filename, file_path

def extract_text_from_pdf(file_path):
    text = ''
    doc = fitz.open(file_path)

    for page in doc:
        text += page.get_text()
    doc.close()
    return text

def extract_text_from_docx(file_path):
    text = ''
    doc = Document(file_path)

    for paragraph in doc.paragraphs:
        text += paragraph.text + '\n'
    return text

def extract_text_from_txt(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()

def extract_text_from_csv(file_path):
    lines = []
    with open(file_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f)
        for row in reader:
            lines.append(', '.join(row))
    return '\n'.join(lines)

def extract_text_from_xlsx(file_path):
    lines = []
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    for row in sheet.iter_rows(values_only=True):
        line = ', '.join([str(cell) for cell in row if cell is not None])
        if line:
            lines.append(line)
        return '\n'.join(lines)

def extract_text(file_path):
    extension = file_path.rsplit('.', 1)[1].lower()
    extractors = {
        'pdf': extract_text_from_pdf,
        'docx': extract_text_from_docx,
        'txt': extract_text_from_txt,
        'csv': extract_text_from_csv,
        'xlsx': extract_text_from_xlsx,
    }
    extractor = extractors.get(extension)
    if extractor:
        return extractor(file_path)

    return ''